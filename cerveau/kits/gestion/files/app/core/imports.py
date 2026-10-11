"""Import de données dans un module depuis un fichier CSV ou Excel (.xlsx).

Tout ou rien : si une ligne est fausse, rien n'est importé et chaque erreur est signalée avec
son numéro de ligne. Les colonnes sont reconnues par le nom ou le libellé du champ, sans tenir
compte des accents ni de la casse. Un lien (client…) se donne par son nom ou son numéro.
"""

from __future__ import annotations

import csv
import io
import unicodedata
from datetime import date, datetime

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.modeles import numero_suivant

LIGNES_MAX = 5000
TAILLE_MAX = 5_000_000


def _cle(texte) -> str:
    texte = unicodedata.normalize("NFKD", str(texte or "").strip().lower())
    return "".join(c for c in texte if c.isalnum())


def lire_fichier(nom: str, contenu: bytes) -> list[list]:
    """Lignes du fichier, en-tête compris."""
    if len(contenu) > TAILLE_MAX:
        raise ValueError("fichier trop gros (5 Mo au plus)")
    if nom.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ValueError(f"fichier Excel illisible ({exc})") from None
        return [list(ligne) for ligne in classeur.active.iter_rows(values_only=True)]
    if not nom.lower().endswith((".csv", ".txt")):
        raise ValueError("format non reconnu : CSV ou Excel (.xlsx)")
    for encodage in ("utf-8-sig", "cp1252"):
        try:
            texte = contenu.decode(encodage)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialecte = csv.Sniffer().sniff(texte[:5000], delimiters=";,\t")
    except csv.Error:
        dialecte = csv.excel
        dialecte.delimiter = ";"
    return [ligne for ligne in csv.reader(io.StringIO(texte), dialecte)]


def _brut(valeur) -> str:
    if valeur is None:
        return ""
    if isinstance(valeur, datetime):
        return valeur.date().isoformat()
    if isinstance(valeur, date):
        return valeur.isoformat()
    if isinstance(valeur, float) and valeur.is_integer():
        return str(int(valeur))
    return str(valeur).strip()


def importer(module, db: Session, lignes: list[list]) -> tuple[list, list[str], list[str]]:
    """Objets prêts à enregistrer, erreurs, colonnes ignorées. N'écrit rien en base."""
    from app.core.framework import REGISTRE, convertir

    lignes = [l for l in lignes if any(_brut(v) for v in l)]
    if len(lignes) < 2:
        return [], ["le fichier ne contient aucune ligne de données"], []
    if len(lignes) - 1 > LIGNES_MAX:
        return [], [f"{len(lignes) - 1} lignes : {LIGNES_MAX} au plus par import"], []
    champs = {}
    for champ in module.champs:
        if champ.type != "calcul":
            champs[_cle(champ.nom)] = champ
            champs[_cle(champ.libelle)] = champ
    entete = [_cle(v) for v in lignes[0]]
    colonnes = {i: champs[c] for i, c in enumerate(entete) if c in champs}
    ignorees = [_brut(lignes[0][i]) for i, c in enumerate(entete) if c and c not in champs]
    erreurs = []
    manquants = [c.libelle for c in module.champs if c.requis and c not in colonnes.values()]
    if manquants:
        return [], [f"colonnes obligatoires absentes : {', '.join(manquants)}"], ignorees

    index_refs: dict[str, dict[str, int]] = {}
    for champ in {c.nom: c for c in colonnes.values()}.values():
        if champ.type == "ref":
            cible = REGISTRE[champ.ref]
            index = {}
            for objet in db.scalars(select(cible.modele)):
                index[_cle(cible.libelle_objet(objet))] = objet.id
                index[str(objet.id)] = objet.id
            index_refs[champ.nom] = index

    objets = []
    for numero, ligne in enumerate(lignes[1:], start=2):
        valeurs = {}
        for i, champ in colonnes.items():
            brut = _brut(ligne[i] if i < len(ligne) else "")
            try:
                if champ.type == "ref":
                    if not brut:
                        valeur = convertir(champ, "")
                    elif brut in index_refs[champ.nom] or _cle(brut) in index_refs[champ.nom]:
                        valeur = index_refs[champ.nom].get(brut) or index_refs[champ.nom][_cle(brut)]
                    else:
                        raise ValueError(f"« {brut} » introuvable dans {REGISTRE[champ.ref].libelle.lower()}")
                elif champ.type == "numero":
                    valeur = brut or None
                else:
                    if champ.type == "date" and "/" in brut:
                        jour, mois, annee = brut.split("/")[:3]
                        brut = f"{int(annee):04d}-{int(mois):02d}-{int(jour):02d}"
                    if champ.type == "choix" and brut:
                        brut = next((c for c in champ.choix if _cle(c) == _cle(brut)), brut)
                    valeur = convertir(champ, brut)
            except ValueError as exc:
                erreurs.append(f"ligne {numero}, {champ.libelle} : {exc}")
                continue
            valeurs[champ.colonne] = valeur
        for champ in module.champs:
            if champ.saisissable and champ.colonne not in valeurs and champ.defaut is not None:
                valeurs[champ.colonne] = champ.valeur_defaut()
        objets.append(valeurs)
    return objets, erreurs, ignorees


def enregistrer(module, db: Session, objets: list[dict]) -> int:
    for valeurs in objets:
        objet = module.modele(**valeurs)
        for champ in module.champs:
            if champ.type == "numero" and not getattr(objet, champ.colonne):
                setattr(objet, champ.colonne, numero_suivant(db, f"{module.nom}.{champ.nom}", champ.format))
        db.add(objet)
    return len(objets)
